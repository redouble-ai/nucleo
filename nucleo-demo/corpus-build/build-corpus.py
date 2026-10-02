#!/usr/bin/env python3
"""Generates the binary half of the demo corpus: Office files, a text PDF, a scanned PDF, images,
and junk. The text half is authored as files in the corpus directory. Everything here is invented
content about a fictional bicycle maker, Halcyon Bicycle Works; nothing is a real company's.

The prices are placed on purpose, so the pricing demo has a history to reconcile: the 2026 list
(effective 1 January 2026) in the spreadsheet, the deck and the invoice; the April 2026 Kestrel
change on the portal screenshot and the whiteboard; a planning figure in the product spec; a
sign that shows what the tune-up used to cost; supplier costs in the parts sheet and on a note.
The dates are the documents' own, where a document carries one.

Usage: build-corpus.py <corpus dir>
Needs: pillow, python-docx, openpyxl, python-pptx, reportlab.
"""
import os
import random
import sys

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

out = sys.argv[1]
os.makedirs(out, exist_ok=True)


def font(size):
    for path in ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Helvetica.ttc",
                 "/Library/Fonts/Arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                pass
    return ImageFont.load_default()


def text_image(lines, width=1100, height=None, size=30, margin=60, leading=44):
    height = height or margin * 2 + leading * len(lines)
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    y = margin
    for line in lines:
        draw.text((margin, y), line, fill="black", font=font(size))
        y += leading
    return image


# --- product-spec.docx: a planning price from before the list was set
doc = Document()
doc.add_heading("Halcyon Meridian 3 - Product Specification", 0)
doc.add_paragraph("Model year 2026. Internal reference HBW-M3. Status: approved for tooling, November 2025.")
doc.add_heading("Frame", 1)
doc.add_paragraph("Double-butted 6061 aluminium, hydroformed top tube, tapered head tube 44 mm to 56 mm. "
                  "Sizes S, M, L, XL. Weight 1,640 g in size M, painted.")
doc.add_heading("Drivetrain", 1)
doc.add_paragraph("1x12, 32T chainring, 10-51T cassette, 12-speed clutch derailleur. Chain line 55 mm.")
doc.add_heading("Wheels and tyres", 1)
doc.add_paragraph("29 inch, 30 mm internal rim width, tubeless ready. Front 2.4 inch, rear 2.35 inch.")
doc.add_heading("Warranty", 1)
doc.add_paragraph("Frame: lifetime for the first owner. Paint and decals: two years. Bearings: one year.")
doc.add_heading("Pricing", 1)
doc.add_paragraph("Target retail price EUR 2,599 (planning figure, November 2025). The list price is set by "
                  "sales in the January price list and may differ.")
table = doc.add_table(rows=1, cols=3)
table.rows[0].cells[0].text = "Size"
table.rows[0].cells[1].text = "Reach (mm)"
table.rows[0].cells[2].text = "Stack (mm)"
for size, reach, stack in (("S", "425", "600"), ("M", "445", "615"), ("L", "465", "630"), ("XL", "485", "645")):
    row = table.add_row().cells
    row[0].text, row[1].text, row[2].text = size, reach, stack
doc.save(os.path.join(out, "product-spec.docx"))

# --- price-list.xlsx: the 2026 list, effective 1 January 2026, and the parts sheet with supplier costs
wb = Workbook()
ws = wb.active
ws.title = "Dealer prices 2026"
ws.append(["Halcyon Bicycle Works - price list 2026, effective from 2026-01-01"])
ws.append(["SKU", "Model", "Size", "Dealer price (EUR)", "Retail price (EUR)", "Stock"])
skus = [("HBW-M3-S", "Meridian 3", "S", 1890, 2699, 12), ("HBW-M3-M", "Meridian 3", "M", 1890, 2699, 31),
        ("HBW-M3-L", "Meridian 3", "L", 1890, 2699, 27), ("HBW-M3-XL", "Meridian 3", "XL", 1890, 2699, 8),
        ("HBW-K1-M", "Kestrel 1 gravel", "M", 1420, 2049, 19), ("HBW-K1-L", "Kestrel 1 gravel", "L", 1420, 2049, 14),
        ("HBW-C2-U", "Comet 2 city", "one size", 690, 999, 55)]
for row in skus:
    ws.append(list(row))
ws2 = wb.create_sheet("Parts")
ws2.append(["Part", "Supplier", "Unit cost (EUR)", "Lead time (weeks)"])
for row in (("Frame M3", "Tai Han Alloy", 310, 14), ("Fork M3", "Tai Han Alloy", 95, 14),
            ("Wheelset 29", "Rimwerk GmbH", 180, 6), ("Drivetrain 1x12", "Kettenbach", 240, 8)):
    ws2.append(list(row))
wb.save(os.path.join(out, "price-list.xlsx"))

# --- dealer-pitch.pptx: January 2026, so its Kestrel price predates the April change
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Halcyon Bicycle Works"
slide.placeholders[1].text = "Dealer programme 2026 - dealer meeting, 12 January 2026"
for title, bullets in (
        ("Why Halcyon", ["Three models, one frame family, 80% shared parts",
                         "Lifetime frame warranty, handled in 5 working days",
                         "Dealer margin 30% at list, 34% on stock orders over 20 units"]),
        ("2026 line-up", ["Meridian 3 trail, from EUR 2,699", "Kestrel 1 gravel, from EUR 2,049",
                          "Comet 2 city, EUR 999"]),
        ("What we ask", ["Two demo bikes on the floor", "A trained mechanic per store",
                         "Warranty claims through the dealer portal, never by email"])):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    body = s.placeholders[1].text_frame
    body.text = bullets[0]
    for b in bullets[1:]:
        p = body.add_paragraph()
        p.text = b
        p.font.size = Pt(20)
prs.save(os.path.join(out, "dealer-pitch.pptx"))

# --- invoice-0417.pdf (text layer): March 2026, at the January dealer prices
c = canvas.Canvas(os.path.join(out, "invoice-0417.pdf"), pagesize=letter)
c.setFont("Helvetica-Bold", 16)
c.drawString(72, 720, "Halcyon Bicycle Works - Invoice 2026-0417")
c.setFont("Helvetica", 11)
lines = ["Bill to: Nordlicht Radsport GmbH, Hafenstrasse 12, 24103 Kiel",
         "Dealer account: D-101        Prices: dealer list 2026",
         "Date: 2026-03-02        Due: 2026-04-01        Terms: net 30",
         "",
         "Qty   SKU          Description                 Unit (EUR)    Total (EUR)",
         "  4   HBW-M3-M     Meridian 3, size M            1,890.00       7,560.00",
         "  2   HBW-K1-L     Kestrel 1 gravel, size L      1,420.00       2,840.00",
         "  6   HBW-C2-U     Comet 2 city                    690.00       4,140.00",
         "",
         "Subtotal                                                       14,540.00",
         "VAT 19%                                                         2,762.60",
         "Total due                                                      17,302.60",
         "",
         "Bank: Sparkasse Holstein, IBAN DE00 0000 0000 0000 0000 00 (specimen)",
         "Thank you for your order. Frames ship from Kiel within 10 working days."]
y = 690
for line in lines:
    c.drawString(72, y, line)
    y -= 18
c.showPage()
c.save()

# --- scanned-letter.pdf (image only, no text layer): a dated letter naming a service price
letter_lines = ["Halcyon Bicycle Works", "Customer Service", "Kiel, 24 February 2026", "",
                "Dear Ms. Aaltonen,", "",
                "Thank you for your letter of 14 February about the",
                "creaking noise from your Meridian 3 headset. This is",
                "covered by the frame warranty. Please bring the bike",
                "to your dealer, who will replace the headset bearings",
                "and check the fork steerer for wear. The bearing",
                "replacement, normally EUR 40, is at no charge.", "",
                "We are sorry for the trouble.", "",
                "Kind regards,", "Jonas Berg, Customer Service"]
page = text_image(letter_lines, width=1200, height=1600, size=32, margin=100, leading=60)
page.save(os.path.join(out, "scanned-letter.pdf"), "PDF", resolution=150.0)

# --- shop-sign.png: a sign with a price and the price it replaced, and no date anywhere
text_image(["HALCYON BICYCLE WORKS", "Workshop hours", "Mon - Fri  9:00 - 18:00", "Sat  10:00 - 14:00",
            "Tune-up EUR 49  (was EUR 45)", "Headset bearing swap EUR 40",
            "Warranty service by appointment", "Ring the bell twice"], width=1000, size=40, leading=64).save(
    os.path.join(out, "shop-sign.png"))

# --- warranty-card.jpg
text_image(["WARRANTY CARD", "Model: Meridian 3    Size: L", "Frame no.: HBW-M3-2026-00841",
            "Purchased: 2026-05-11", "Dealer: Nordlicht Radsport, Kiel",
            "Lifetime frame warranty, first owner only"], width=1000, size=34, leading=56).convert("RGB").save(
    os.path.join(out, "warranty-card.jpg"), "JPEG", quality=88)

# --- a photo with no text at all: a plain gradient
photo = Image.new("RGB", (800, 500))
px = photo.load()
for x in range(800):
    for y in range(500):
        px[x, y] = (int(40 + x / 800 * 120), int(90 + y / 500 * 100), 160)
photo.save(os.path.join(out, "sky.png"))

# --- mixed-report.pdf: page 1 has a text layer, page 2 is a picture of a table (a scanned attachment)
c = canvas.Canvas(os.path.join(out, "mixed-report.pdf"), pagesize=letter)
c.setFont("Helvetica-Bold", 15)
c.drawString(72, 720, "Quarterly quality report, Q1 2026")
c.setFont("Helvetica", 11)
for i, line in enumerate(["Warranty claims received: 44, of which 31 headset creak (see February meeting).",
                          "Each bearing swap costs us about EUR 40 in parts and dealer labour.",
                          "Average answer time: 3.2 working days against the 5-day promise.",
                          "Paint defects on Meridian 3 pre-series: 2 of 40 frames, both orange peel on the",
                          "down tube, traced to booth humidity on 3 March.",
                          "The attached page is the scanned inspection sheet from the paint line."]):
    c.drawString(72, 690 - i * 18, line)
c.showPage()
sheet = text_image(["PAINT LINE INSPECTION SHEET", "Week 10, booth 2",
                    "Frame          Colour        Result", "HBW-M3-2026-00801   Signal orange   pass",
                    "HBW-M3-2026-00802   Signal orange   orange peel, down tube",
                    "HBW-M3-2026-00803   Slate           pass",
                    "HBW-M3-2026-00804   Signal orange   orange peel, down tube",
                    "Booth humidity 71%, above the 60% limit", "Inspector: M. Solheim"],
                   width=1200, height=1600, size=30, margin=100, leading=58)
c.drawImage(ImageReader(sheet), 36, 100, width=540, height=720)
c.showPage()
c.save()

# --- screenshot-prices.png: the portal after the April change, dated on the screen
text_image(["Dealer portal - prices as of 2026-04-12 (EUR)", "",
            "SKU         Model              Size   Dealer   Retail",
            "HBW-M3-M    Meridian 3         M      1,890    2,699",
            "HBW-K1-L    Kestrel 1 gravel   L      1,450    2,099",
            "HBW-C2-U    Comet 2 city       one    690      999"], width=1200, size=28, leading=48).save(
    os.path.join(out, "screenshot-prices.png"))

# --- whiteboard.jpg: notes photographed off a whiteboard, no date on it
board = Image.new("RGB", (1200, 900), (232, 240, 236))
draw = ImageDraw.Draw(board)
for i, line in enumerate(["Launch dealers - call by Friday", "Paint line wk 16 - Mira confirms",
                          "Headset claims: 31 -> swap bearings", "Kestrel 1: 2049 -> 2099 from 1 Apr (Anke)",
                          "2027 terms: Comet out of stock margin", "Next meeting 4 March, 9:00"]):
    draw.text((80, 80 + i * 120), line, fill=(30, 50, 140), font=font(44))
board.save(os.path.join(out, "whiteboard.jpg"), "JPEG", quality=85)

# --- notes.txt that is not text: a JPEG under a text extension, with a supplier's price on it
text_image(["Handwritten-ish note", "Order 60 derailleurs, ETA 20 March",
            "Kettenbach: derailleur now 52 EUR each (was 48)", "Ask about the clutch spring"],
           width=1000, size=32, leading=56).convert("RGB").save(os.path.join(out, "notes.txt"), "JPEG", quality=85)

# --- binaries the classifier must call: random bytes under two extensions
random.seed(20260914)
with open(os.path.join(out, "telemetry.dat"), "wb") as f:
    f.write(bytes(random.getrandbits(8) for _ in range(4096)))
with open(os.path.join(out, "keys.enc"), "wb") as f:
    f.write(b"\x00\x01ENC\x02" + bytes(random.getrandbits(8) for _ in range(2048)))
print("corpus binaries written to", out)
